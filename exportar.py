"""
exportar.py
Genera un archivo Excel con el análisis de simulaciones PVsyst,
replicando el formato del Ejemplo.xlsx (una hoja por proyecto).

Uso:
    python exportar.py                      # guarda como 'reporte_simulaciones.xlsx'
    python exportar.py mi_reporte.xlsx      # nombre personalizado
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from db_manager import get_all_simulations, get_monthly, init_db

# --- Constantes ---------------------------------------------------------------

OUTPUT_DEFAULT = Path("reporte_simulaciones.xlsx")

MESES_ES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

# Columnas: B=2 (etiquetas), C-N=3-14 (meses 1-12), O=15 (anual)
COL_LABEL   = 2
COL_MES_INI = 3
COL_MES_FIN = 14
COL_ANUAL   = 15

# Filas del layout (igual que Ejemplo.xlsx)
ROW_PROYECTO   = 4
ROW_SOLARGIS   = 5
ROW_P90_DC     = 6
ROW_P90_DIARIA = 7
ROW_RATIO      = 8
ROW_N_MES      = 11
ROW_MES        = 12
ROW_PROPORCION = 13
ROW_P50        = 14
ROW_P75        = 15
ROW_P90        = 16
ROW_P99        = 17


# --- Generador de hoja --------------------------------------------------------

def _escribir_hoja(wb: Workbook, sim: dict, monthly: list[dict]) -> None:
    """Crea y rellena una hoja con el layout del Ejemplo.xlsx."""

    nombre_hoja = (sim.get("project_code") or sim["filename"])[:31]
    ws = wb.create_sheet(title=nombre_hoja)

    # Valores calculados reutilizados
    p50_kwh   = (sim.get("p50_annual_energy_mwh") or 0) * 1000
    p90_kwh   = (sim.get("p90_annual_energy_mwh") or 0) * 1000
    p99_kwh   = (sim.get("p99_annual_energy_mwh") or 0) * 1000
    p75_kwh   = (sim.get("p75_annual_energy_mwh") or 0) * 1000
    p75_est   = bool(sim.get("p75_is_estimated", 1))   # True -> mostrar 0 (no disponible en PDF)
    solargis  = sim.get("solargis_daily_irradiation")
    ratio     = sim.get("ratio_yield_irradiation")

    monthly_sorted = sorted(monthly, key=lambda r: r["month_number"])

    # -- Bloque de encabezado --------------------------------------------------
    ws.cell(ROW_PROYECTO,   COL_MES_INI).value = sim.get("project_code") or sim["filename"]

    ws.cell(ROW_SOLARGIS,   COL_LABEL).value   = "Radiación Día Solargis"
    ws.cell(ROW_SOLARGIS,   COL_MES_INI).value = solargis

    ws.cell(ROW_P90_DC,     COL_LABEL).value   = "Producción en P90 en DC"
    ws.cell(ROW_P90_DC,     COL_MES_INI).value = p90_kwh or None

    ws.cell(ROW_P90_DIARIA, COL_LABEL).value   = "Producción en P90 diaria"
    ws.cell(ROW_P90_DIARIA, COL_MES_INI).value = sim.get("p90_daily_yield")

    ws.cell(ROW_RATIO,      COL_LABEL).value   = "Ratio"
    ws.cell(ROW_RATIO,      COL_MES_INI).value = ratio

    # -- Fila de números de mes ------------------------------------------------
    ws.cell(ROW_N_MES, COL_LABEL).value = "# del Mes"
    for col, num in zip(range(COL_MES_INI, COL_MES_FIN + 1), range(1, 13)):
        ws.cell(ROW_N_MES, col).value = num
    ws.cell(ROW_N_MES, COL_ANUAL).value = "-"

    # -- Fila de nombres de mes ------------------------------------------------
    ws.cell(ROW_MES, COL_LABEL).value = "Mes"
    for col, mes in zip(range(COL_MES_INI, COL_MES_FIN + 1), MESES_ES):
        ws.cell(ROW_MES, col).value = mes
    ws.cell(ROW_MES, COL_ANUAL).value = "Año"

    # -- Fila de proporciones --------------------------------------------------
    ws.cell(ROW_PROPORCION, COL_LABEL).value = "Proporción mensual"
    for col, m in zip(range(COL_MES_INI, COL_MES_FIN + 1), monthly_sorted):
        ws.cell(ROW_PROPORCION, col).value = m.get("monthly_proportion")
    ws.cell(ROW_PROPORCION, COL_ANUAL).value = 1

    # -- Fila P50 -------------------------------------------------------------
    ws.cell(ROW_P50, COL_LABEL).value = "P50"
    for col, m in zip(range(COL_MES_INI, COL_MES_FIN + 1), monthly_sorted):
        ws.cell(ROW_P50, col).value = m.get("p50_kwh")
    ws.cell(ROW_P50, COL_ANUAL).value = p50_kwh or None

    # -- Fila P75 (0 cuando no vino en el PDF) --------------------------------
    ws.cell(ROW_P75, COL_LABEL).value = "P75"
    if p75_est:
        for col in range(COL_MES_INI, COL_ANUAL + 1):
            ws.cell(ROW_P75, col).value = 0
    else:
        for col, m in zip(range(COL_MES_INI, COL_MES_FIN + 1), monthly_sorted):
            ws.cell(ROW_P75, col).value = m.get("p75_kwh")
        ws.cell(ROW_P75, COL_ANUAL).value = p75_kwh or None

    # -- Fila P90 -------------------------------------------------------------
    ws.cell(ROW_P90, COL_LABEL).value = "P90"
    for col, m in zip(range(COL_MES_INI, COL_MES_FIN + 1), monthly_sorted):
        ws.cell(ROW_P90, col).value = m.get("p90_kwh")
    ws.cell(ROW_P90, COL_ANUAL).value = p90_kwh or None

    # -- Fila P99 -------------------------------------------------------------
    ws.cell(ROW_P99, COL_LABEL).value = "P99"
    for col, m in zip(range(COL_MES_INI, COL_MES_FIN + 1), monthly_sorted):
        ws.cell(ROW_P99, col).value = m.get("p99_kwh")
    ws.cell(ROW_P99, COL_ANUAL).value = p99_kwh or None

    # -- Formato ---------------------------------------------------------------
    negrita = Font(name="Arial", size=11, bold=True)
    normal  = Font(name="Arial", size=11)
    centro  = Alignment(horizontal="center")

    # Ancho de columnas
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions[get_column_letter(COL_LABEL)].width = 28
    for c in range(COL_MES_INI, COL_ANUAL + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    # Negrita en etiquetas
    filas_label = [
        ROW_SOLARGIS, ROW_P90_DC, ROW_P90_DIARIA, ROW_RATIO,
        ROW_N_MES, ROW_MES, ROW_PROPORCION,
        ROW_P50, ROW_P75, ROW_P90, ROW_P99,
    ]
    for fila in filas_label:
        ws.cell(fila, COL_LABEL).font = negrita

    # Centrar y fuente normal en toda la zona de datos
    for fila in range(ROW_N_MES, ROW_P99 + 1):
        for col in range(COL_MES_INI, COL_ANUAL + 1):
            cell = ws.cell(fila, col)
            cell.font  = normal
            cell.alignment = centro


# --- Exportador ---------------------------------------------------------------

def export(output_path: Path) -> None:
    init_db()
    sims = get_all_simulations()

    if not sims:
        print("[!] No hay simulaciones en la base de datos.")
        print("    Ejecutá primero: python main.py\n")
        return

    wb = Workbook()
    wb.remove(wb.active)  # eliminar hoja vacía por defecto

    for sim in sims:
        monthly = get_monthly(sim["id"])
        _escribir_hoja(wb, sim, monthly)
        codigo = sim.get("project_code") or sim["filename"]
        solargis_ok = "OK" if sim.get("solargis_daily_irradiation") else "SIN Solargis"
        print(f"  OK {codigo}  ({solargis_ok})")

    wb.save(output_path)
    print(f"\nOK Excel guardado en: {output_path.resolve()}\n")


# --- Punto de entrada ----------------------------------------------------------

if __name__ == "__main__":
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT_DEFAULT
    print(f"\nGenerando reporte Excel -> '{output}' ...\n")
    export(output)
